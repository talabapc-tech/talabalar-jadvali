# -*- coding: utf-8 -*-
"""
01-guruh (26-01) qat'iy "Farmatsiya ishi", qolgan barcha guruhlar qat'iy "Hamshiralik ishi"
qilib bazalarni (Excel, students.json, talabalar_bazasi.json) bir xillashtirish skripti.
"""

import os
import sys
import json
import shutil
import time
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, 'data', 'Talabalar_Toliq_Royxati.xlsx')
STUDENTS_PATH = os.path.join(BASE_DIR, 'data', 'students.json')
BAZA_PATH = os.path.join(BASE_DIR, 'data', 'talabalar_bazasi.json')
BACKUP_DIR = os.path.join(BASE_DIR, 'data', 'backups')
os.makedirs(BACKUP_DIR, exist_ok=True)

stamp = time.strftime('%Y%m%d_%H%M%S')

# 1. Zaxira nusxa olish
shutil.copy2(EXCEL_PATH, os.path.join(BACKUP_DIR, f"Talabalar_Toliq_Royxati_pre_yon_{stamp}.xlsx"))
shutil.copy2(STUDENTS_PATH, os.path.join(BACKUP_DIR, f"students_pre_yon_{stamp}.json"))
shutil.copy2(BAZA_PATH, os.path.join(BACKUP_DIR, f"talabalar_bazasi_pre_yon_{stamp}.json"))
print(f"✅ Zaxira nusxalar olindi: data/backups/*_{stamp}.*")

# 2. students.json yangilash
with open(STUDENTS_PATH, 'r', encoding='utf-8') as f:
    students = json.load(f)

count_farm = 0
count_hamsh = 0

for s in students:
    grp = str(s.get('group') or '').strip()
    if grp == '26-01' or grp.endswith('-01'):
        s['yon'] = 'Farmatsiya ishi'
        count_farm += 1
    else:
        s['yon'] = 'Hamshiralik ishi'
        count_hamsh += 1

with open(STUDENTS_PATH, 'w', encoding='utf-8') as f:
    json.dump(students, f, ensure_ascii=False, indent=2)

print(f"✅ students.json yangilandi: Farmatsiya ishi = {count_farm}, Hamshiralik ishi = {count_hamsh}")

# 3. talabalar_bazasi.json yangilash
with open(BAZA_PATH, 'r', encoding='utf-8') as f:
    baza_obj = json.load(f)

b_students = baza_obj.get('students', [])
for s in b_students:
    grp = str(s.get('group') or '').strip()
    if grp == '26-01' or grp.endswith('-01'):
        s['yon'] = 'Farmatsiya ishi'
    else:
        s['yon'] = 'Hamshiralik ishi'

baza_obj['updated_at'] = time.strftime('%Y-%m-%d %H:%M:%S')
with open(BAZA_PATH, 'w', encoding='utf-8') as f:
    json.dump(baza_obj, f, ensure_ascii=False, indent=2)

print(f"✅ talabalar_bazasi.json yangilandi ({len(b_students)} ta talaba).")

# 4. Talabalar_Toliq_Royxati.xlsx yangilash
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active

excel_farm = 0
excel_hamsh = 0

for r in range(2, ws.max_row + 1):
    fish = str(ws.cell(r, 8).value or ws.cell(r, 2).value or '').strip()
    if not fish:
        continue
    grp = str(ws.cell(r, 23).value or '').strip()
    if grp == '26-01' or grp.endswith('-01'):
        ws.cell(r, 3).value = 'Farmatsiya ishi'
        excel_farm += 1
    else:
        ws.cell(r, 3).value = 'Hamshiralik ishi'
        excel_hamsh += 1

wb.save(EXCEL_PATH)
print(f"✅ Talabalar_Toliq_Royxati.xlsx yangilandi: Farmatsiya ishi = {excel_farm}, Hamshiralik ishi = {excel_hamsh}")

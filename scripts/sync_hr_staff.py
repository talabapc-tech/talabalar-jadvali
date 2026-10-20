# -*- coding: utf-8 -*-
r"""
SH.T.T XODIMLAR VA O'QITUVCHILAR BAZASINI SINXRONLASH VA EXCEL EXPORT QILISH
==========================================================================
Manba papka: D:\Ozodbek HR\ASOSIY XUJJATLAR\Shahrisabz\1.Barcha xodimlar\SH.T.T
1. 01. Shtat.xlsx       -> ShTAT - xodimlar (8 nafar)
2. 02. Asosiy.xlsx      -> Asosiy o'qituvchilar (13 nafar)
3. 03. O'rindosh.xlsx   -> O'rindosh o'qituvchilar (8 nafar)
4. BARCHA_O'QITUVCHILAR_MALUMOTLARI.xlsx -> qo'shimcha yashash joyi, pasport berilgan joy
5. 000. SH.T.T BAZA.xlsx                 -> qo'shimcha OTM va mutaxassislik

Chiqish fayllari (Hech qanday tagchiziq '_' ishlatilmaydi, faqat probel):
- data/xodimlar.json
- hisobotlar/Oqituvchilar royxati.xlsx
- hisobotlar/Shtat xodimlar royxati.xlsx
- hisobotlar/Barcha xodimlar royxati.xlsx
"""

import os
import sys
import json
import re
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    if hasattr(sys.stderr, 'reconfigure'):
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HR_PATH = r"D:\Ozodbek HR\ASOSIY XUJJATLAR\Shahrisabz\1.Barcha xodimlar\SH.T.T"
OUTPUT_JSON = os.path.join(BASE_DIR, 'data', 'xodimlar.json')
HISOBOTLAR_DIR = os.path.join(BASE_DIR, 'hisobotlar')

def clean_val(v):
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%d.%m.%Y')
    return str(v).strip()

def clean_phone(tel):
    if not tel:
        return ''
    digits = re.sub(r'\D', '', str(tel))
    if len(digits) == 9:
        return f'+998 {digits[:2]} {digits[2:5]} {digits[5:7]} {digits[7:]}'
    if len(digits) == 12 and digits.startswith('998'):
        return f'+{digits[:3]} {digits[3:5]} {digits[5:8]} {digits[8:10]} {digits[10:]}'
    return str(tel).strip()

def sync_staff_from_hr(hr_dir=HR_PATH):
    if not os.path.exists(hr_dir):
        print(f"[XODIMLAR] ⚠️ HR manzil topilmadi: {hr_dir}")
        return []

    # 1. Qo'shimcha manbalarni o'qish (yashash joyi, pasport berilgan joy, OTM)
    extra_teachers = {}
    f_all_teachers = os.path.join(hr_dir, "BARCHA_O'QITUVCHILAR_MALUMOTLARI.xlsx")
    if os.path.exists(f_all_teachers):
        try:
            wb = openpyxl.load_workbook(f_all_teachers, data_only=True)
            ws = wb.active
            for r in range(3, ws.max_row + 1):
                fish = clean_val(ws.cell(r, 2).value)
                if not fish:
                    continue
                key = re.sub(r'\s+', ' ', fish.lower()).strip()
                extra_teachers[key] = {
                    'pasport_berilgan_joy': clean_val(ws.cell(r, 8).value),
                    'yashash_joyi': clean_val(ws.cell(r, 9).value),
                }
        except Exception as e:
            print(f"[XODIMLAR] BARCHA_O'QITUVCHILAR o'qishda xato: {e}")

    extra_baza = {}
    f_baza = os.path.join(hr_dir, '000. SH.T.T BAZA.xlsx')
    if os.path.exists(f_baza):
        try:
            wb = openpyxl.load_workbook(f_baza, data_only=True)
            ws = wb.active
            for r in range(4, ws.max_row + 1):
                fish = clean_val(ws.cell(r, 2).value)
                if not fish:
                    continue
                key = re.sub(r'\s+', ' ', fish.lower()).strip()
                extra_baza[key] = {
                    'tamomlagan_otm': clean_val(ws.cell(r, 9).value),
                    'otm_tamom_yili': clean_val(ws.cell(r, 10).value),
                    'mutaxassislik': clean_val(ws.cell(r, 11).value),
                }
        except Exception as e:
            print(f"[XODIMLAR] 000. SH.T.T BAZA o'qishda xato: {e}")

    staff_list = []

    # 2. 01. Shtat.xlsx -> ShTAT - xodimlar
    f_shtat = os.path.join(hr_dir, '01. Shtat.xlsx')
    if os.path.exists(f_shtat):
        wb1 = openpyxl.load_workbook(f_shtat, data_only=True)
        ws1 = wb1.active
        for r in range(4, ws1.max_row + 1):
            num = clean_val(ws1.cell(r, 2).value)
            fish = clean_val(ws1.cell(r, 4).value)
            if not fish:
                continue
            norm_key = re.sub(r'\s+', ' ', fish.lower()).strip()
            baza_info = extra_baza.get(norm_key, {})
            lavozim = clean_val(ws1.cell(r, 3).value)
            
            # Bo'lim aniqlash
            if any(k in lavozim.lower() for k in ['direktor']):
                bolim = 'Rahbariyat'
            elif any(k in lavozim.lower() for k in ['kadr', "o'ibdo", "o‘ibdo", 'sotuv']):
                bolim = "Ma'muriy"
            elif any(k in lavozim.lower() for k in ['qorovul', 'farrosh', "xo'jalik"]):
                bolim = "Xo'jalik"
            elif 'hamshira' in lavozim.lower():
                bolim = "Tibbiy xizmat"
            else:
                bolim = "Boshqa"

            staff_list.append({
                'id': f'X-SHTAT-{int(num):02d}' if num.isdigit() else f'X-SHTAT-{r-3:02d}',
                'kategoriya': 'ShTAT - xodimlar',
                'ish_turi': 'Shtat',
                'bolim': bolim,
                'lavozim': lavozim,
                'fish': fish,
                'tel': clean_phone(ws1.cell(r, 5).value),
                'dob': clean_val(ws1.cell(r, 6).value),
                'tugilgan_joy': clean_val(ws1.cell(r, 7).value),
                'pv': clean_val(ws1.cell(r, 8).value).upper(),
                'pinfl': clean_val(ws1.cell(r, 9).value),
                'tamomlagan_otm': baza_info.get('tamomlagan_otm', ''),
                'otm_tamom_yili': baza_info.get('otm_tamom_yili', ''),
                'mutaxassislik': baza_info.get('mutaxassislik', ''),
                'buyruq_sana': clean_val(ws1.cell(r, 10).value),
                'ishga_sana': clean_val(ws1.cell(r, 10).value),
                'stavka': '1.0',
                'status': 'faol',
                'izoh': ''
            })

    # 3. 02. Asosiy.xlsx -> Asosiy o'qituvchilar
    f_asosiy = os.path.join(hr_dir, '02. Asosiy.xlsx')
    if os.path.exists(f_asosiy):
        wb2 = openpyxl.load_workbook(f_asosiy, data_only=True)
        ws2 = wb2.active
        for r in range(4, ws2.max_row + 1):
            num = clean_val(ws2.cell(r, 2).value)
            fish = clean_val(ws2.cell(r, 3).value)
            if not fish:
                continue
            norm_key = re.sub(r'\s+', ' ', fish.lower()).strip()
            ext = extra_teachers.get(norm_key, {})
            
            staff_list.append({
                'id': f'X-ASOSIY-{int(num):02d}' if num.isdigit() else f'X-ASOSIY-{r-3:02d}',
                'kategoriya': 'Asosiy o\'qituvchilar',
                'ish_turi': 'Asosiy',
                'bolim': 'Pedagog',
                'lavozim': 'Asosiy o\'qituvchi',
                'fish': fish,
                'tel': clean_phone(ws2.cell(r, 4).value),
                'dob': clean_val(ws2.cell(r, 5).value),
                'tugilgan_joy': clean_val(ws2.cell(r, 6).value),
                'pv': clean_val(ws2.cell(r, 7).value).upper(),
                'pinfl': clean_val(ws2.cell(r, 8).value),
                'tamomlagan_otm': clean_val(ws2.cell(r, 9).value),
                'otm_tamom_yili': clean_val(ws2.cell(r, 10).value),
                'mutaxassislik': clean_val(ws2.cell(r, 11).value),
                'buyruq_sana': clean_val(ws2.cell(r, 12).value),
                'ishga_sana': clean_val(ws2.cell(r, 12).value),
                'pasport_berilgan_joy': ext.get('pasport_berilgan_joy', ''),
                'yashash_joyi': ext.get('yashash_joyi', ''),
                'stavka': '1.0',
                'status': 'faol',
                'izoh': ''
            })

    # 4. 03. O'rindosh.xlsx -> O'rindosh o'qituvchilar
    f_orindosh = os.path.join(hr_dir, "03. O'rindosh.xlsx")
    if os.path.exists(f_orindosh):
        wb3 = openpyxl.load_workbook(f_orindosh, data_only=True)
        ws3 = wb3.active
        for r in range(4, ws3.max_row + 1):
            num = clean_val(ws3.cell(r, 2).value)
            fish = clean_val(ws3.cell(r, 3).value)
            if not fish:
                continue
            norm_key = re.sub(r'\s+', ' ', fish.lower()).strip()
            ext = extra_teachers.get(norm_key, {})
            
            staff_list.append({
                'id': f'X-ORINDOSH-{int(num):02d}' if num.isdigit() else f'X-ORINDOSH-{r-3:02d}',
                'kategoriya': 'O\'rindosh o\'qituvchilar',
                'ish_turi': 'O\'rindosh',
                'bolim': 'Pedagog',
                'lavozim': 'O\'rindosh o\'qituvchi',
                'fish': fish,
                'tel': clean_phone(ws3.cell(r, 4).value),
                'dob': clean_val(ws3.cell(r, 5).value),
                'tugilgan_joy': clean_val(ws3.cell(r, 6).value),
                'pv': clean_val(ws3.cell(r, 7).value).upper(),
                'pinfl': clean_val(ws3.cell(r, 8).value),
                'tamomlagan_otm': clean_val(ws3.cell(r, 9).value),
                'otm_tamom_yili': clean_val(ws3.cell(r, 10).value),
                'mutaxassislik': clean_val(ws3.cell(r, 11).value),
                'buyruq_sana': clean_val(ws3.cell(r, 12).value),
                'ishga_sana': clean_val(ws3.cell(r, 12).value),
                'pasport_berilgan_joy': ext.get('pasport_berilgan_joy', ''),
                'yashash_joyi': ext.get('yashash_joyi', ''),
                'stavka': '0.5',
                'status': 'faol',
                'izoh': ''
            })

    # 5. data/xodimlar.json ga saqlash
    os.makedirs(os.path.dirname(OUTPUT_JSON), exist_ok=True)
    with open(OUTPUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(staff_list, f, ensure_ascii=False, indent=2)

    print(f"[XODIMLAR] ✅ {len(staff_list)} nafar xodim data/xodimlar.json ga saqlandi.")
    return staff_list

def export_custom_excel(records, filename_no_underscore, title, selected_columns=None):
    """
    Excel eksport funksiyasi:
    - Fayl nomida hech qachon '_' (tagchiziq) ishlatilmaydi, faqat bo'sh joy (probel)!
    - Foydalanuvchi tanlagan ustunlar bo'yicha tuziladi.
    """
    os.makedirs(HISOBOTLAR_DIR, exist_ok=True)
    
    # Fayl nomidagi '_' larni probelga almashtirish
    clean_filename = filename_no_underscore.replace('_', ' ')
    if not clean_filename.endswith('.xlsx'):
        clean_filename += '.xlsx'
    
    out_path = os.path.join(HISOBOTLAR_DIR, clean_filename)

    ALL_COLUMNS = {
        'num': ('№', 6),
        'fish': ('F.I.SH (Familiya, Ism, Sharif)', 32),
        'lavozim': ('Lavozimi', 22),
        'kategoriya': ('Kategoriya', 20),
        'ish_turi': ('Ish turi', 14),
        'bolim': ("Bo'limi", 16),
        'tel': ('Telefon raqami', 18),
        'dob': ("Tug'ilgan sanasi", 15),
        'tugilgan_joy': ("Tug'ilgan joyi", 30),
        'pv': ('Pasport seriya va raqami', 15),
        'pinfl': ('JSHSHIR (PINFL)', 18),
        'tamomlagan_otm': ("Tamomlagan OTM / Ta'lim muassasasi", 34),
        'otm_tamom_yili': ('OTM tamomlagan yili', 16),
        'mutaxassislik': ("Diplom bo'yicha mutaxassisligi", 28),
        'buyruq_sana': ('Ishga kirgan buyruq sanasi', 18),
        'yashash_joyi': ('Doimiy yashash manzili', 30),
        'pasport_berilgan_joy': ('Pasport berilgan joyi', 26),
        'stavka': ('Stavka', 10),
        'maosh': ('Oylik maosh', 15),
        'status': ('Holati', 12),
        'izoh': ('Izoh', 22),
    }

    if not selected_columns:
        selected_columns = ['num', 'fish', 'lavozim', 'tel', 'dob', 'tugilgan_joy', 'pv', 'pinfl', 'tamomlagan_otm', 'mutaxassislik', 'buyruq_sana']

    cols_to_use = [c for c in selected_columns if c in ALL_COLUMNS]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]

    # Sarlavha blok
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols_to_use))
    title_cell = ws.cell(1, 1, f"SHAHRISABZ TIBBIYOT TEXNIKUMI — {title.upper()}")
    title_cell.font = Font(name='Arial', size=13, bold=True, color='107C41')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    # Ustun sarlavhalari (Row 2)
    header_fill = PatternFill(start_color='107C41', end_color='107C41', fill_type='solid')
    header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    border_side = Side(border_style='thin', color='D9D9D9')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center')

    for col_idx, col_key in enumerate(cols_to_use, 1):
        label, width = ALL_COLUMNS[col_key]
        c = ws.cell(2, col_idx, label)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center_align
        c.border = cell_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[2].height = 26

    # Ma'lumotlarni yozish
    row_font = Font(name='Arial', size=10)
    for r_idx, item in enumerate(records, 3):
        ws.row_dimensions[r_idx].height = 20
        for c_idx, col_key in enumerate(cols_to_use, 1):
            if col_key == 'num':
                val = r_idx - 2
            elif col_key == 'status':
                st = item.get('status')
                val = "Ta'tilda" if st == 'tatilda' else ("Bo'shatilgan" if st == 'boshatilgan' else "Faol")
            else:
                val = item.get(col_key, '')
            
            cell = ws.cell(r_idx, c_idx, val)
            cell.font = row_font
            cell.border = cell_border
            if col_key in ('num', 'dob', 'pv', 'pinfl', 'otm_tamom_yili', 'buyruq_sana', 'stavka', 'status'):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    wb.save(out_path)
    print(f"[EXCEL] ✅ Yaratildi: {out_path} ({len(records)} ta yozuv, {len(cols_to_use)} ta ustun)")
    return out_path

def generate_default_staff_reports():
    """Standart 3 ta kategoriya bo'yicha Excel hisobotlarni yaratish (spaces only)"""
    if not os.path.exists(OUTPUT_JSON):
        sync_staff_from_hr()

    with open(OUTPUT_JSON, 'r', encoding='utf-8') as f:
        staff = json.load(f)

    # 1. Barcha o'qituvchilar (Asosiy + O'rindosh)
    teachers = [s for s in staff if "o'qituvchi" in s.get('kategoriya', '').lower() or s.get('bolim') == 'Pedagog']
    export_custom_excel(
        teachers,
        "Oqituvchilar royxati.xlsx",
        "O'qituvchilar ro'yxati",
        ['num', 'fish', 'kategoriya', 'tel', 'dob', 'tugilgan_joy', 'pv', 'pinfl', 'tamomlagan_otm', 'otm_tamom_yili', 'mutaxassislik', 'buyruq_sana']
    )

    # 2. Shtat xodimlari
    shtat = [s for s in staff if s.get('ish_turi') == 'Shtat' or 'shtat' in s.get('kategoriya', '').lower()]
    export_custom_excel(
        shtat,
        "Shtat xodimlar royxati.xlsx",
        "ShTAT xodimlar ro'yxati",
        ['num', 'lavozim', 'fish', 'tel', 'dob', 'tugilgan_joy', 'pv', 'pinfl', 'tamomlagan_otm', 'mutaxassislik', 'buyruq_sana']
    )

    # 3. Barcha xodimlar
    export_custom_excel(
        staff,
        "Barcha xodimlar royxati.xlsx",
        "Barcha xodimlar ro'yxati",
        ['num', 'fish', 'kategoriya', 'lavozim', 'bolim', 'tel', 'dob', 'pv', 'pinfl', 'tamomlagan_otm', 'mutaxassislik', 'buyruq_sana', 'status']
    )

if __name__ == '__main__':
    records = sync_staff_from_hr()
    generate_default_staff_reports()
